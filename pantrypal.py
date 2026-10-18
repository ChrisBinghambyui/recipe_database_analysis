"""
Pantry Pal
Finds recipes you can make with what's in your pantry. Uses the RecipeNLG
dataset from Kaggle:
https://www.kaggle.com/datasets/wilmerarltstrmberg/recipe-dataset-over-2m

Setup: put recipes_data.csv in the same folder as this file. If it's somewhere
else, the program asks once and remembers where it is.

What it does:
  - Search recipes by title, must contain, and cannot contain
  - Pantry list you can edit and save
  - "Use Pantry" sorts recipes by how much you're missing (can make now,
    only seasonings missing, buy 1 item, buy 2 items)
  - Best Buys tab shows which single ingredient would unlock the most recipes
  - Favorites and custom lists (a recipe can be in more than one list)
  - Open recipes in their own window and export to .md or .docx

To run: pip install pandas python-docx, then python pantry_pal.py
(python-docx is only needed for the Word export)
"""

import json
import queue
import re
import threading
import webbrowser
from collections import Counter
from functools import lru_cache
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk

import pandas as pd

# ----
# Settings
# ----
BASE_DIR = Path(__file__).resolve().parent
DATASET_NAMES = ("recipes_data.csv", "full_dataset.csv")  # names to look for
PANTRY_PATH = BASE_DIR / "pantry.json"
LIBRARY_PATH = BASE_DIR / "library.json"    # favorites and lists
SETTINGS_PATH = BASE_DIR / "settings.json"  # remembers where the csv is

# The full csv is ~2.2 million recipes and eats a ton of RAM, so we only load
# the first chunk. Bump this up (or set to None) if your computer can take it.
MAX_ROWS = 200_000

RESULT_CAP = 500             # max rows shown per tab so the window stays fast
STAPLES = ("water", "salt")  # can be assumed to be in everyone's kitchen
FAVORITES = "Favorites"      # built-in list, can't be renamed or deleted

# words that don't usually matter when comparing ingredients
FILLER_WORDS = {
    "fresh", "large", "small", "medium", "chopped", "diced", "minced",
    "sliced", "frozen", "optional", "finely", "freshly", "whole",
}

# seasonings are easy to borrow, substitute, or skip
SEASONING_WORDS = {
    "salt", "pepper", "paprika", "cumin", "oregano", "basil", "thyme",
    "rosemary", "sage", "parsley", "cinnamon", "nutmeg", "cayenne",
    "turmeric", "coriander", "cardamom", "allspice", "dill", "bay",
    "seasoning", "spice", "herb", "curry", "tarragon", "marjoram",
    "saffron", "mace", "vanilla", "chive",
}
# "bell pepper" has "pepper" in it but it isn't a seasoning
NOT_SEASONING_WORDS = {"bell", "jalapeno", "sauce", "sweet", "stuffed", "roasted"}


# ----
# Text helpers (cleaning up the ingredient names so they can be compared)
# ----
def singular(word):
    """Basic plural remover so 'tomatoes' and 'tomato' match."""
    if len(word) > 3 and word.endswith("ies"):
        return word[:-3] + "y"
    if len(word) > 3 and word.endswith(("ches", "shes", "sses", "xes", "oes")):
        return word[:-2]
    if len(word) > 3 and word.endswith("s") and not word.endswith("ss"):
        return word[:-1]
    return word


@lru_cache(maxsize=None)
def normalize(name):
    """'Onions, chopped' -> 'onion'. Cached since it runs a LOT."""
    words = re.findall(r"[a-z]+", name.lower())
    words = [singular(w) for w in words if w not in FILLER_WORDS]
    return " ".join(words)


@lru_cache(maxsize=None)
def is_seasoning(key):
    """Guesses if an ingredient is a spice or herb."""
    words = set(key.split())
    return bool(words & SEASONING_WORDS) and not (words & NOT_SEASONING_WORDS)


def parse_list(text):
    """The csv stores lists as JSON text. Returns [] if the row is broken."""
    try:
        value = json.loads(text)
    except (ValueError, TypeError):
        return []
    return value if isinstance(value, list) else []


def build_url(link):
    """
    Makes a dataset link actually open in a browser. The csv has no https://
    on its links, and cookbooks.com breaks if you keep the www. part because it's older, so that
    gets removed for that site.
    """
    link = link.strip()
    if not link:
        return ""
    if link.startswith("http"):
        return link
    if link.startswith("www.cookbooks.com"):
        link = link[len("www."):]
    return "https://" + link


def split_terms(text):
    """'chicken, garlic ,  ' -> ['chicken', 'garlic']"""
    return [t.strip() for t in text.split(",") if t.strip()]


def safe_filename(name):
    """Removes characters that aren't allowed in file names."""
    return re.sub(r"[^\w\- ]+", "", name).strip()[:60] or "recipe"


# ----
# Settings file (so you only have to pick the csv once)
# ----
def load_settings():
    try:
        return json.loads(SETTINGS_PATH.read_text(encoding="utf-8"))
    except (ValueError, OSError):
        return {}


def save_settings(settings):
    SETTINGS_PATH.write_text(json.dumps(settings, indent=2), encoding="utf-8")


def find_dataset():
    """Tries the saved path first, then looks for the csv by name."""
    saved = load_settings().get("dataset_path")
    if saved and Path(saved).exists():
        return Path(saved)
    for folder in (BASE_DIR, Path.cwd()):
        for name in DATASET_NAMES:
            candidate = folder / name
            if candidate.exists():
                return candidate
    return None


# ----
# Export stuff (Markdown and Word)
# A "recipe" below is a dict with: id, title, link, ingredients, directions
# ----
def recipe_markdown(recipe, level=1):
    """Turns one recipe into Markdown. level = how many # the title gets."""
    h = "#" * level
    lines = [f"{h} {recipe['title']}", ""]
    url = build_url(recipe["link"])
    if url:
        lines += [f"Source: <{url}>", ""]
    lines += [f"{h}# Ingredients", ""]
    lines += [f"- {item}" for item in recipe["ingredients"]]
    lines += ["", f"{h}# Directions", ""]
    lines += [f"{n}. {step}" for n, step in enumerate(recipe["directions"], start=1)]
    return "\n".join(lines) + "\n"


def list_markdown(name, recipes):
    """Whole list as one Markdown file: list name on top, recipes under it."""
    parts = [f"# {name}\n"] + [recipe_markdown(r, level=2) for r in recipes]
    return "\n".join(parts)


def write_docx(path, heading, recipes, whole_list):
    """
    Saves recipes to a Word file using python-docx. The import is down here
    so the rest of the program still works if it isn't installed.
    """
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("Word export needs the python-docx package.\n"
                           "Install it with:  pip install python-docx")
    doc = Document()
    if whole_list:
        doc.add_heading(heading, 0)
    recipe_level = 1 if whole_list else 0  # single recipe uses the Title style
    for recipe in recipes:
        doc.add_heading(recipe["title"], recipe_level)
        url = build_url(recipe["link"])
        if url:
            doc.add_paragraph(f"Source: {url}")
        doc.add_heading("Ingredients", recipe_level + 1)
        for item in recipe["ingredients"]:
            doc.add_paragraph(item, style="List Bullet")
        doc.add_heading("Directions", recipe_level + 1)
        # numbering the steps by hand, Word's numbered list doesn't restart
        # between recipes
        for n, step in enumerate(recipe["directions"], start=1):
            doc.add_paragraph(f"{n}. {step}")
    doc.save(path)


# ----
# Library (favorites and custom lists, saved in library.json)
# ----
class Library:
    """
    Each list only stores recipe ids. The actual recipe data is kept in one
    separate dict, so a saved recipe still works even if it isn't in the rows
    currently loaded from the csv.
    """

    def __init__(self, path):
        self.path = path
        self.lists = {FAVORITES: []}  # list name -> [recipe ids]
        self.recipes = {}             # recipe id -> recipe dict
        self.load()

    def load(self):
        if not self.path.exists():
            return
        try:
            data = json.loads(self.path.read_text(encoding="utf-8"))
        except (ValueError, OSError):
            return
        saved_lists = data.get("lists", {})
        self.recipes = data.get("recipes", {})
        self.lists = {FAVORITES: list(saved_lists.get(FAVORITES, []))}  # keep Favorites first
        for name, ids in saved_lists.items():
            if name != FAVORITES:
                self.lists[name] = list(ids)

    def save(self):
        data = {"lists": self.lists, "recipes": self.recipes}
        self.path.write_text(json.dumps(data, indent=2), encoding="utf-8")

    def contains(self, name, rid):
        return rid in self.lists.get(name, [])

    def add(self, name, recipe):
        self.recipes[recipe["id"]] = recipe
        if recipe["id"] not in self.lists[name]:
            self.lists[name].append(recipe["id"])

    def remove(self, name, rid):
        if rid in self.lists[name]:
            self.lists[name].remove(rid)
        self.forget_if_orphaned(rid)

    def forget_if_orphaned(self, rid):
        """If no list uses the recipe anymore, delete its saved copy."""
        if not any(rid in ids for ids in self.lists.values()):
            self.recipes.pop(rid, None)

    def create_list(self, name):
        if name in self.lists:
            return False
        self.lists[name] = []
        return True

    def rename_list(self, old, new):
        if old == FAVORITES or new in self.lists:
            return False
        self.lists = {(new if k == old else k): v for k, v in self.lists.items()}
        return True

    def delete_list(self, name):
        if name == FAVORITES:
            return
        for rid in self.lists.pop(name, []):
            self.forget_if_orphaned(rid)

    def recipes_in(self, name):
        return [self.recipes[rid] for rid in self.lists[name] if rid in self.recipes]


# ----
# Recipe data and matching (no GUI stuff in this part)
# ----
class RecipeStore:
    """Holds all the recipes plus an index so ingredient lookups are fast."""

    def __init__(self):
        self.titles = []
        self.titles_lower = []
        self.keys = []             # for each recipe: tuple of cleaned ingredient names
        self.raw_ingredients = []  # JSON text, only parsed when a recipe is shown
        self.raw_directions = []
        self.links = []
        self.vocab = {}            # ingredient name -> set of its words
        self.token_index = {}      # word -> ingredient names that contain it

    def load(self, path, max_rows):
        """Reads the csv with pandas, cleans it up, and builds the lookups."""
        df = pd.read_csv(
            path,
            usecols=["title", "ingredients", "directions", "link", "NER"],
            nrows=max_rows,
        )
        # no title or no ingredient names = not useful, drop it
        df = df.dropna(subset=["title", "NER"])

        for row in df.itertuples(index=False):
            names = parse_list(row.NER)
            keys = tuple(dict.fromkeys(  # removes duplicates but keeps the order
                k for k in (normalize(n) for n in names if isinstance(n, str)) if k
            ))
            if len(keys) < 2:  # skip 1 ingredient "recipes", they flood the results
                continue
            self.titles.append(str(row.title))
            self.titles_lower.append(str(row.title).lower())
            self.keys.append(keys)
            self.raw_ingredients.append(row.ingredients if isinstance(row.ingredients, str) else "[]")
            self.raw_directions.append(row.directions if isinstance(row.directions, str) else "[]")
            self.links.append(row.link if isinstance(row.link, str) else "")
            for key in keys:
                if key not in self.vocab:
                    words = frozenset(key.split())
                    self.vocab[key] = words
                    for word in words:
                        self.token_index.setdefault(word, set()).add(key)

    def recipe_id(self, i):
        """Id for a recipe: the link, or the title if there's no link."""
        return self.links[i] or self.titles[i]

    def recipe_at(self, i):
        """Builds the full recipe dict for one row (this is where the JSON gets parsed)."""
        link = self.links[i]
        return {
            "id": self.recipe_id(i),  # links are unique per recipe
            "title": self.titles[i],
            "link": link,
            "ingredients": [str(x) for x in parse_list(self.raw_ingredients[i])],
            "directions": [str(x) for x in parse_list(self.raw_directions[i])],
        }

    def keys_matching(self, term):
        """Every ingredient containing all words of the term ('chicken' finds 'chicken breast')."""
        words = frozenset(normalize(term).split())
        if not words:
            return set()
        sets = sorted((self.token_index.get(w, set()) for w in words), key=len)
        return set.intersection(*sets) if sets[0] else set()

    def search(self, title, must, exclude):
        """Returns the indexes of recipes that pass the title/must/exclude filters."""
        if not (title or must or exclude):
            return list(range(len(self.keys)))
        must_sets = [self.keys_matching(t) for t in must]
        if any(not s for s in must_sets):  # nobody uses a required ingredient
            return []
        banned = set()
        for term in exclude:
            banned |= self.keys_matching(term)

        hits = []
        for i, keys in enumerate(self.keys):
            if title and title not in self.titles_lower[i]:
                continue
            if banned and not banned.isdisjoint(keys):
                continue
            if all(not s.isdisjoint(keys) for s in must_sets):
                hits.append(i)
        return hits

    def covered_keys(self, pantry_items, loose):
        """
        Returns every ingredient the pantry covers.
        Strict: all the words of a recipe ingredient have to be in one pantry
                item ('olive oil' covers 'oil' and 'olive oil').
        Loose:  also lets a pantry item cover longer ingredients that contain
                it ('chicken' covers 'chicken breast').
        """
        covered = set()
        for item in pantry_items:
            words = frozenset(normalize(item).split())
            if not words:
                continue
            candidates = set()
            for w in words:
                candidates |= self.token_index.get(w, set())
            covered.update(k for k in candidates if self.vocab[k] <= words)
            if loose:
                covered |= self.keys_matching(item)
        return covered


# ----
# GUI
# ----
class PantryWindow(tk.Toplevel):
    """Little window for editing the pantry list."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("My Pantry")
        self.geometry("320x420")

        entry_row = ttk.Frame(self)
        entry_row.pack(fill="x", padx=8, pady=8)
        self.entry = ttk.Entry(entry_row)
        self.entry.pack(side="left", fill="x", expand=True)
        self.entry.bind("<Return>", lambda e: self.add())
        ttk.Button(entry_row, text="Add", command=self.add).pack(side="left", padx=(6, 0))

        ttk.Label(self, text="Tip: add several at once, separated by commas.").pack(anchor="w", padx=8)

        list_row = ttk.Frame(self)
        list_row.pack(fill="both", expand=True, padx=8, pady=8)
        self.listbox = tk.Listbox(list_row, selectmode="extended")
        scroll = ttk.Scrollbar(list_row, command=self.listbox.yview)
        self.listbox.config(yscrollcommand=scroll.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        scroll.pack(side="left", fill="y")
        self.listbox.bind("<Delete>", lambda e: self.remove())

        ttk.Button(self, text="Remove selected", command=self.remove).pack(pady=(0, 8))
        self.refresh()
        self.entry.focus_set()

    def refresh(self):
        self.listbox.delete(0, "end")
        for item in self.app.pantry:
            self.listbox.insert("end", item)

    def add(self):
        for term in split_terms(self.entry.get().lower()):
            if term not in self.app.pantry:
                self.app.pantry.append(term)
        self.app.pantry.sort()
        self.entry.delete(0, "end")
        self.app.save_pantry()
        self.refresh()

    def remove(self):
        doomed = {self.listbox.get(i) for i in self.listbox.curselection()}
        self.app.pantry = [p for p in self.app.pantry if p not in doomed]
        self.app.save_pantry()
        self.refresh()


class ListPicker(tk.Toplevel):
    """
    The "Save to list" popup. Checking or unchecking a list adds or removes
    the recipe right away. You can also type a name to make a new list.
    """

    def __init__(self, app, recipe):
        super().__init__(app)
        self.app = app
        self.recipe = recipe
        self.vars = []  # have to keep these around or tkinter throws them out
        self.title("Save to list")
        self.transient(app)
        self.resizable(False, False)

        ttk.Label(self, text=recipe["title"], wraplength=300,
                  font=("TkDefaultFont", 10, "bold")).pack(anchor="w", padx=12, pady=(12, 6))
        self.body = ttk.Frame(self)
        self.body.pack(fill="both", expand=True, padx=12)

        new_row = ttk.Frame(self)
        new_row.pack(fill="x", padx=12, pady=10)
        self.new_entry = ttk.Entry(new_row)
        self.new_entry.pack(side="left", fill="x", expand=True)
        self.new_entry.bind("<Return>", lambda e: self.create())
        ttk.Button(new_row, text="New list", command=self.create).pack(side="left", padx=(6, 0))
        ttk.Button(self, text="Done", command=self.destroy).pack(pady=(0, 12))
        self.rebuild()
        self.new_entry.focus_set()

    def rebuild(self):
        """Redraws the checkboxes (checked = recipe is already in that list)."""
        for widget in self.body.winfo_children():
            widget.destroy()
        self.vars = []
        library = self.app.library
        for name in library.lists:
            var = tk.BooleanVar(value=library.contains(name, self.recipe["id"]))
            self.vars.append(var)
            ttk.Checkbutton(self.body, text=name, variable=var,
                            command=lambda n=name, v=var: self.toggle(n, v)).pack(anchor="w", pady=1)

    def toggle(self, name, var):
        if var.get():
            self.app.library.add(name, self.recipe)
        else:
            self.app.library.remove(name, self.recipe["id"])
        self.app.library.save()
        self.app.library_changed()

    def create(self):
        name = self.new_entry.get().strip()
        if not name:
            return
        if not self.app.library.create_list(name):
            messagebox.showinfo("List exists", f"You already have a list called '{name}'.", parent=self)
            return
        self.app.library.add(name, self.recipe)
        self.app.library.save()
        self.new_entry.delete(0, "end")
        self.rebuild()
        self.app.library_changed()


class RecipeView(ttk.Frame):
    """
    Shows a recipe with its buttons (favorite, save to list, pop out, export).
    Used in the bottom of the main window and in the pop-out windows.
    """

    def __init__(self, parent, app, popout=True):
        super().__init__(parent)
        self.app = app
        self.recipe = None
        self.url = ""

        bar = ttk.Frame(self)
        bar.pack(fill="x", pady=(0, 4))
        self.fav_btn = ttk.Button(bar, text="\u2606 Favorite", command=self.toggle_favorite)
        self.list_btn = ttk.Button(bar, text="Save to list...", command=self.save_to_list)
        self.fav_btn.pack(side="left")
        self.list_btn.pack(side="left", padx=4)
        self.controls = [self.fav_btn, self.list_btn]
        if popout:
            window_btn = ttk.Button(bar, text="Open in window", command=self.pop_out)
            window_btn.pack(side="left")
            self.controls.append(window_btn)
        self.export_btn = ttk.Menubutton(bar, text="Export")
        menu = tk.Menu(self.export_btn, tearoff=False)
        menu.add_command(label="Markdown (.md)", command=lambda: self.export("md"))
        menu.add_command(label="Word document (.docx)", command=lambda: self.export("docx"))
        self.export_btn["menu"] = menu
        self.export_btn.pack(side="left", padx=4)
        self.controls.append(self.export_btn)

        body = ttk.Frame(self)
        body.pack(fill="both", expand=True)
        self.text = tk.Text(body, wrap="word", height=12, state="disabled")
        scroll = ttk.Scrollbar(body, command=self.text.yview)
        self.text.config(yscrollcommand=scroll.set)
        self.text.pack(side="left", fill="both", expand=True)
        scroll.pack(side="left", fill="y")
        self.text.tag_config("heading", font=("TkDefaultFont", 13, "bold"))
        self.text.tag_config("section", font=("TkDefaultFont", 10, "bold"))
        self.text.tag_config("missing", foreground="#b00020")
        self.text.tag_config("link", foreground="#1a5fb4", underline=True)
        self.text.tag_bind("link", "<Button-1>", lambda e: self.open_link())
        self.text.tag_bind("link", "<Enter>", lambda e: self.text.config(cursor="hand2"))
        self.text.tag_bind("link", "<Leave>", lambda e: self.text.config(cursor=""))

        self.refresh_buttons()

    def show(self, recipe, missing=()):
        """Displays a recipe. missing = pantry gaps to highlight in red."""
        self.recipe = recipe
        box = self.text
        box.config(state="normal")
        box.delete("1.0", "end")
        box.insert("end", recipe["title"] + "\n", "heading")

        self.url = build_url(recipe["link"])
        if self.url:
            box.insert("end", self.url + "\n", "link")
        if missing:
            box.insert("end", "Missing from pantry: " + ", ".join(missing) + "\n", "missing")

        box.insert("end", "\nIngredients\n", "section")
        for line in recipe["ingredients"]:
            box.insert("end", f"  - {line}\n")
        box.insert("end", "\nDirections\n", "section")
        for step, line in enumerate(recipe["directions"], start=1):
            box.insert("end", f"  {step}. {line}\n")
        box.config(state="disabled")
        self.refresh_buttons()

    def refresh_buttons(self):
        """Buttons are disabled until a recipe is showing. Also updates the star."""
        state = "normal" if self.recipe else "disabled"
        for control in self.controls:
            control.config(state=state)
        if self.recipe and self.app.library.contains(FAVORITES, self.recipe["id"]):
            self.fav_btn.config(text="\u2605 Favorited")
        else:
            self.fav_btn.config(text="\u2606 Favorite")

    def toggle_favorite(self):
        library, recipe = self.app.library, self.recipe
        if library.contains(FAVORITES, recipe["id"]):
            library.remove(FAVORITES, recipe["id"])
        else:
            library.add(FAVORITES, recipe)
        library.save()
        self.app.library_changed()

    def save_to_list(self):
        ListPicker(self.app, self.recipe)

    def pop_out(self):
        self.app.open_recipe_window(self.recipe, self.current_missing())

    def current_missing(self):
        return self.app.last_missing.get(self.recipe["id"], [])

    def export(self, kind):
        self.app.export([self.recipe], self.recipe["title"], kind, whole_list=False)

    def open_link(self):
        if self.url:
            webbrowser.open(self.url)


class RecipeWindow(tk.Toplevel):
    """A recipe in its own window. You can have a bunch open at once."""

    def __init__(self, app, recipe, missing):
        super().__init__(app)
        self.title(recipe["title"])
        self.geometry("640x720")
        view = RecipeView(self, app, popout=False)
        view.pack(fill="both", expand=True, padx=8, pady=8)
        view.show(recipe, missing)
        app.views.append(view)  # keeps the favorite star synced with the main window


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Pantry Pal")
        self.geometry("1050x800")

        self.store = RecipeStore()
        self.library = Library(LIBRARY_PATH)
        self.ready = False
        self.jobs = queue.Queue()      # background threads put their results here
        self.pantry = self.load_pantry()
        self.pantry_window = None
        self.views = []                # all the RecipeViews that need updating
        self.last_missing = {}         # recipe id -> what's missing from the pantry

        self.assume_staples = tk.BooleanVar(value=True)
        self.loose_matching = tk.BooleanVar(value=False)
        self.status = tk.StringVar(value="Starting up...")

        self.build_filters()
        self.build_results()
        ttk.Label(self, textvariable=self.status, relief="sunken", anchor="w").pack(fill="x", side="bottom")

        self.after(100, self.poll_jobs)
        self.start_loading()

    # ----- layout -----
    def build_filters(self):
        box = ttk.LabelFrame(self, text="Filters (apply to both Search and Use Pantry)")
        box.pack(fill="x", padx=8, pady=8)
        box.columnconfigure(1, weight=1)

        self.title_var, self.must_var, self.exclude_var = tk.StringVar(), tk.StringVar(), tk.StringVar()
        fields = [
            ("Title contains", self.title_var),
            ("Must contain (comma separated)", self.must_var),
            ("Cannot contain (comma separated)", self.exclude_var),
        ]
        for row, (label, var) in enumerate(fields):
            ttk.Label(box, text=label).grid(row=row, column=0, sticky="w", padx=6, pady=3)
            entry = ttk.Entry(box, textvariable=var)
            entry.grid(row=row, column=1, sticky="ew", padx=6, pady=3)
            entry.bind("<Return>", lambda e: self.on_search())

        options = ttk.Frame(box)
        options.grid(row=0, column=2, rowspan=3, padx=10)
        ttk.Checkbutton(options, text="Assume water and salt", variable=self.assume_staples).pack(anchor="w")
        ttk.Checkbutton(options, text="Loose matching (chicken = chicken breast)",
                        variable=self.loose_matching).pack(anchor="w")

        buttons = ttk.Frame(box)
        buttons.grid(row=3, column=0, columnspan=3, pady=6)
        self.search_btn = ttk.Button(buttons, text="Search", command=self.on_search)
        self.pantry_btn = ttk.Button(buttons, text="Use Pantry", command=self.on_use_pantry)
        self.search_btn.pack(side="left", padx=4)
        self.pantry_btn.pack(side="left", padx=4)
        ttk.Button(buttons, text="Edit Pantry", command=self.open_pantry).pack(side="left", padx=4)

    def build_results(self):
        pane = ttk.PanedWindow(self, orient="vertical")
        pane.pack(fill="both", expand=True, padx=8, pady=(0, 8))

        self.notebook = ttk.Notebook(pane)
        pane.add(self.notebook, weight=3)

        recipe_cols = [("title", "Recipe", 430, False), ("ingredients", "Ingredients", 80, True),
                       ("missing", "Missing", 70, True), ("items", "Missing items", 330, False)]
        self.tabs = {}  # key -> (frame, tree, tab title)
        self.add_tab("search", "Search Results", [recipe_cols[0], recipe_cols[1]])
        self.add_tab("can", "Can Make Now", recipe_cols)
        self.add_tab("season", "Only Seasonings Missing", recipe_cols)
        self.add_tab("buy1", "Buy 1 Item", recipe_cols)
        self.add_tab("buy2", "Buy 2 Items", recipe_cols)
        self.add_tab("best", "Best Buys", [("title", "Ingredient to buy", 430, False),
                                           ("ingredients", "Recipes it unlocks", 140, True)])
        self.build_lists_tab()

        # click a row to see the recipe, double click to open it in a window
        for key, (_, tree, _) in self.tabs.items():
            if key != "best":
                tree.bind("<<TreeviewSelect>>", lambda e, t=tree: self.on_select(t))
                tree.bind("<Double-1>", lambda e, t=tree: self.on_double_click(t))

        self.view = RecipeView(pane, self)
        pane.add(self.view, weight=2)
        self.views.append(self.view)

    def add_tab(self, key, text, columns):
        frame = ttk.Frame(self.notebook)
        tree = self.make_tree(frame, columns)
        self.notebook.add(frame, text=text)
        self.tabs[key] = (frame, tree, text)

    def make_tree(self, parent, columns):
        """Makes a table with a scrollbar. Clicking a heading sorts it."""
        tree = ttk.Treeview(parent, columns=[c[0] for c in columns], show="headings", selectmode="browse")
        for col, heading, width, numeric in columns:
            tree.heading(col, text=heading, command=lambda c=col, t=tree: self.sort_by(t, c, False))
            tree.column(col, width=width, anchor="center" if numeric else "w")
        scroll = ttk.Scrollbar(parent, command=tree.yview)
        tree.config(yscrollcommand=scroll.set)
        tree.pack(side="left", fill="both", expand=True)
        scroll.pack(side="left", fill="y")
        return tree

    def build_lists_tab(self):
        """My Lists tab: list names on the left, recipes in that list on the right."""
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text="My Lists")

        left = ttk.Frame(frame)
        left.pack(side="left", fill="y", padx=(6, 6), pady=6)
        self.list_box = tk.Listbox(left, exportselection=False, width=26)
        self.list_box.pack(fill="both", expand=True)
        self.list_box.bind("<<ListboxSelect>>", lambda e: self.show_selected_list())
        self.list_names = []  # same order as the rows in the listbox

        buttons = ttk.Frame(left)
        buttons.pack(fill="x", pady=(6, 0))
        ttk.Button(buttons, text="New", width=7, command=self.new_list).pack(side="left")
        ttk.Button(buttons, text="Rename", width=7, command=self.rename_list).pack(side="left", padx=3)
        ttk.Button(buttons, text="Delete", width=7, command=self.delete_list).pack(side="left")
        export_btn = ttk.Menubutton(left, text="Export whole list")
        export_menu = tk.Menu(export_btn, tearoff=False)
        export_menu.add_command(label="Markdown (.md)", command=lambda: self.export_current_list("md"))
        export_menu.add_command(label="Word document (.docx)", command=lambda: self.export_current_list("docx"))
        export_btn["menu"] = export_menu
        export_btn.pack(fill="x", pady=(6, 0))

        right = ttk.Frame(frame)
        right.pack(side="left", fill="both", expand=True, pady=6)
        table = ttk.Frame(right)
        table.pack(fill="both", expand=True)
        self.lists_tree = self.make_tree(table, [("title", "Recipe", 480, False),
                                                 ("ingredients", "Ingredients", 90, True)])
        self.lists_tree.bind("<<TreeviewSelect>>", lambda e: self.on_select_saved())
        self.lists_tree.bind("<Double-1>", lambda e: self.on_double_click(self.lists_tree, saved=True))
        ttk.Button(right, text="Remove from this list", command=self.remove_from_list).pack(anchor="e", pady=(6, 0))
        self.refresh_lists_tab()

    # ----- small utilities ------
    def load_pantry(self):
        if PANTRY_PATH.exists():
            try:
                return sorted(json.loads(PANTRY_PATH.read_text()))
            except (ValueError, OSError):
                pass
        return []

    def save_pantry(self):
        PANTRY_PATH.write_text(json.dumps(self.pantry, indent=2))
        self.status.set(f"Pantry saved ({len(self.pantry)} items).")

    def open_pantry(self):
        if self.pantry_window and self.pantry_window.winfo_exists():
            self.pantry_window.lift()
        else:
            self.pantry_window = PantryWindow(self)

    def sort_by(self, tree, col, reverse):
        """Sorts a table by a column. Click again to flip the order."""
        rows = [(tree.set(i, col), i) for i in tree.get_children("")]
        try:
            rows.sort(key=lambda r: float(r[0]), reverse=reverse)
        except ValueError:
            rows.sort(key=lambda r: r[0].lower(), reverse=reverse)
        for position, (_, iid) in enumerate(rows):
            tree.move(iid, "", position)
        tree.heading(col, command=lambda: self.sort_by(tree, col, not reverse))

    def fill(self, key, rows, total=None):
        """Replaces a tab's rows. Each row looks like (iid, value1, value2, ...)."""
        frame, tree, base = self.tabs[key]
        tree.delete(*tree.get_children())
        for row in rows[:RESULT_CAP]:
            tree.insert("", "end", iid=str(row[0]), values=row[1:])
        shown_total = len(rows) if total is None else total
        self.notebook.tab(frame, text=f"{base} ({shown_total:,})")

    # ----- background work ------
    def run_async(self, work, done):
        """Runs work in a thread so the window doesn't freeze, then calls done on the main thread."""
        def target():
            try:
                result = work()
            except Exception as exc:  # pass the error back instead of crashing the thread
                result = exc
            self.jobs.put((done, result))
        threading.Thread(target=target, daemon=True).start()

    def poll_jobs(self):
        try:
            while True:
                done, result = self.jobs.get_nowait()
                done(result)
        except queue.Empty:
            pass
        self.after(100, self.poll_jobs)

    def start_loading(self):
        path = find_dataset()  # only asks if it can't find the csv
        if path is None:
            picked = filedialog.askopenfilename(title="Select recipes_data.csv",
                                                filetypes=[("CSV files", "*.csv")])
            if not picked:
                self.status.set("No dataset selected. Restart and choose the recipe CSV.")
                return
            path = Path(picked)
            settings = load_settings()
            settings["dataset_path"] = str(path)  # save it so it doesn't ask again
            save_settings(settings)
        self.set_busy(True)
        self.status.set(f"Loading up to {MAX_ROWS or 'all'} recipes from {path.name}... this can take a minute.")
        self.run_async(lambda: self.store.load(path, MAX_ROWS), self.on_loaded)

    def on_loaded(self, result):
        self.set_busy(False)
        if isinstance(result, Exception):
            messagebox.showerror("Could not load dataset", str(result))
            self.status.set("Dataset failed to load.")
            return
        self.ready = True
        self.status.set(f"Ready: {len(self.store.keys):,} recipes, "
                        f"{len(self.store.vocab):,} distinct ingredients.")

    def set_busy(self, busy):
        state = "disabled" if busy else "normal"
        self.search_btn.config(state=state)
        self.pantry_btn.config(state=state)

    # ----- searching --------------
    def read_filters(self):
        return (self.title_var.get().strip().lower(),
                split_terms(self.must_var.get()),
                split_terms(self.exclude_var.get()))

    def on_search(self):
        if not self.ready:
            return
        title, must, exclude = self.read_filters()
        self.set_busy(True)
        self.status.set("Searching...")
        self.run_async(lambda: self.store.search(title, must, exclude), self.on_search_done)

    def on_search_done(self, hits):
        self.set_busy(False)
        if isinstance(hits, Exception):
            self.status.set(f"Search failed: {hits}")
            return
        rows = [(i, self.store.titles[i], len(self.store.keys[i])) for i in hits]
        self.fill("search", rows)
        self.notebook.select(self.tabs["search"][0])
        self.status.set(f"{len(hits):,} matching recipes (showing the first {RESULT_CAP}).")

    def on_use_pantry(self):
        if not self.ready:
            return
        if not self.pantry:
            messagebox.showinfo("Empty pantry", "Add a few ingredients with Edit Pantry first.")
            return
        title, must, exclude = self.read_filters()
        items = list(self.pantry) + (list(STAPLES) if self.assume_staples.get() else [])
        loose = self.loose_matching.get()
        self.set_busy(True)
        self.status.set("Checking every recipe against your pantry...")
        self.run_async(lambda: self.scan_pantry(title, must, exclude, items, loose), self.on_scan_done)

    def scan_pantry(self, title, must, exclude, items, loose):
        """Sorts recipes into buckets based on what's missing. Runs in a background thread."""
        store = self.store
        covered = store.covered_keys(items, loose)
        buckets = {"can": [], "season": [], "buy1": [], "buy2": []}
        unlock_counts = Counter()  # ingredient -> how many "Buy 1" recipes it would finish
        missing_lookup = {}

        for i in store.search(title, must, exclude):
            keys = store.keys[i]
            missing = [k for k in keys if k not in covered]
            if len(missing) > 5:  # too much missing, skip it
                continue
            seasonings = [k for k in missing if is_seasoning(k)]
            others = [k for k in missing if not is_seasoning(k)]

            if not missing:
                bucket = "can"
            elif not others:
                bucket = "season"
            elif len(others) == 1 and len(seasonings) <= 2:
                bucket = "buy1"
                unlock_counts[others[0]] += 1
            elif len(others) == 2 and len(seasonings) <= 2:
                bucket = "buy2"
            else:
                continue
            buckets[bucket].append((i, len(keys), len(missing)))
            missing_lookup[i] = missing

        # least missing first, then recipes with more ingredients
        for rows in buckets.values():
            rows.sort(key=lambda r: (r[2], -r[1]))
        return buckets, unlock_counts, missing_lookup

    def on_scan_done(self, result):
        self.set_busy(False)
        if isinstance(result, Exception):
            self.status.set(f"Pantry search failed: {result}")
            return
        buckets, unlock_counts, by_index = result
        # store by recipe id so saved recipes and pop-out windows can find it too
        self.last_missing = {self.store.recipe_id(i): m for i, m in by_index.items()}
        for key, rows in buckets.items():
            table = [(i, self.store.titles[i], n, m, ", ".join(by_index[i])) for i, n, m in rows]
            self.fill(key, table)
        best = [(n, name, count) for n, (name, count) in enumerate(unlock_counts.most_common(100))]
        self.fill("best", best)
        self.notebook.select(self.tabs["can"][0])
        self.status.set(f"Pantry scan done: {len(buckets['can']):,} recipes you can make right now.")

    # ----- showing and opening recipes ----------
    def on_select(self, tree):
        selection = tree.selection()
        if selection:
            recipe = self.store.recipe_at(int(selection[0]))
            self.view.show(recipe, self.last_missing.get(recipe["id"], []))

    def on_select_saved(self):
        selection = self.lists_tree.selection()
        recipe = self.library.recipes.get(selection[0]) if selection else None
        if recipe:
            self.view.show(recipe, self.last_missing.get(recipe["id"], []))

    def on_double_click(self, tree, saved=False):
        selection = tree.selection()
        if not selection:
            return
        recipe = self.library.recipes.get(selection[0]) if saved else self.store.recipe_at(int(selection[0]))
        if recipe:
            self.open_recipe_window(recipe, self.last_missing.get(recipe["id"], []))

    def open_recipe_window(self, recipe, missing):
        RecipeWindow(self, recipe, missing)

    # ----- favorites and lists ------------------
    def library_changed(self):
        """Call this after changing a list so all the open views and tabs update."""
        self.views = [v for v in self.views if v.winfo_exists()]
        for view in self.views:
            view.refresh_buttons()
        self.refresh_lists_tab()

    def selected_list(self):
        selection = self.list_box.curselection()
        return self.list_names[selection[0]] if selection else None

    def refresh_lists_tab(self):
        previous = self.selected_list() or FAVORITES
        self.list_names = list(self.library.lists)
        self.list_box.delete(0, "end")
        for name in self.list_names:
            self.list_box.insert("end", f"{name} ({len(self.library.lists[name])})")
        if previous in self.list_names:
            self.list_box.selection_set(self.list_names.index(previous))
        self.show_selected_list()

    def show_selected_list(self):
        tree = self.lists_tree
        tree.delete(*tree.get_children())
        name = self.selected_list()
        if name is None:
            return
        for recipe in self.library.recipes_in(name):
            tree.insert("", "end", iid=recipe["id"], values=(recipe["title"], len(recipe["ingredients"])))

    def new_list(self):
        name = simpledialog.askstring("New list", "Name for the new list:", parent=self)
        if name and name.strip():
            if not self.library.create_list(name.strip()):
                messagebox.showinfo("List exists", f"You already have a list called '{name.strip()}'.")
                return
            self.library.save()
            self.library_changed()

    def rename_list(self):
        old = self.selected_list()
        if old is None or old == FAVORITES:
            messagebox.showinfo("Rename", "Pick one of your own lists first. Favorites keeps its name.")
            return
        new = simpledialog.askstring("Rename list", "New name:", initialvalue=old, parent=self)
        if new and new.strip() and new.strip() != old:
            if not self.library.rename_list(old, new.strip()):
                messagebox.showinfo("Rename", f"A list called '{new.strip()}' already exists.")
                return
            self.library.save()
            self.refresh_lists_tab()
            self.list_box.selection_clear(0, "end")
            self.list_box.selection_set(self.list_names.index(new.strip()))
            self.show_selected_list()

    def delete_list(self):
        name = self.selected_list()
        if name is None or name == FAVORITES:
            messagebox.showinfo("Delete", "Pick one of your own lists first. Favorites cannot be deleted.")
            return
        if messagebox.askyesno("Delete list", f"Delete '{name}'? The recipes themselves stay in any other lists."):
            self.library.delete_list(name)
            self.library.save()
            self.library_changed()

    def remove_from_list(self):
        name, selection = self.selected_list(), self.lists_tree.selection()
        if name and selection:
            self.library.remove(name, selection[0])
            self.library.save()
            self.library_changed()

    # ----- exporting ----------
    def export_current_list(self, kind):
        name = self.selected_list()
        if name is None:
            messagebox.showinfo("Export", "Pick a list first.")
            return
        self.export(self.library.recipes_in(name), name, kind, whole_list=True)

    def export(self, recipes, heading, kind, whole_list):
        """Asks where to save, then writes the recipes as Markdown ('md') or Word ('docx')."""
        if not recipes:
            messagebox.showinfo("Export", "There is nothing to export yet.")
            return
        ext = ".md" if kind == "md" else ".docx"
        label = "Markdown file" if kind == "md" else "Word document"
        path = filedialog.asksaveasfilename(
            title="Export recipe" + ("s" if whole_list else ""),
            defaultextension=ext,
            initialfile=safe_filename(heading) + ext,
            filetypes=[(label, "*" + ext)],
        )
        if not path:
            return
        try:
            if kind == "md":
                text = list_markdown(heading, recipes) if whole_list else recipe_markdown(recipes[0])
                Path(path).write_text(text, encoding="utf-8")
            else:
                write_docx(path, heading, recipes, whole_list)
        except Exception as exc:
            messagebox.showerror("Export failed", str(exc))
            return
        self.status.set(f"Exported to {path}")


if __name__ == "__main__":
    App().mainloop()